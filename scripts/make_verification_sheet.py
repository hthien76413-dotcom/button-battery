#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成逐例人工核对工作表 / Per-case verification worksheet.

把每一例需要判断的病历原文预先拼好，放进一张 Excel。核对者读一行、勾一格，
绝大多数病例不必再回 HIS 调阅——只有存疑的才需要看原始病历。

输出:
  output/核对工作表.xlsx   （含患者层面文本，已 gitignore，不得提交仓库）

填完后转成脚本使用的最终队列:
  python3 scripts/make_verification_sheet.py --to-csv output/核对工作表.xlsx

排序：最可能出问题的排在最前——先是自动标记需核对的，再是部位未定位的、
无影像的，最后才是判定明确的常规病例。

用法:
  python3 scripts/make_verification_sheet.py
  python3 scripts/make_verification_sheet.py --to-csv output/核对工作表.xlsx
"""
import argparse
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import cohort as C

OUTDIR = "output"
OUTFILE = os.path.join(OUTDIR, "核对工作表.xlsx")
SHEET = "核对"
LEGEND = "说明"

FONT = "Arial"
INK = "1A1A1A"
HEAD_FILL = PatternFill("solid", fgColor="E6EAEC")
INPUT_FILL = PatternFill("solid", fgColor="FFF6CC")     # 需填写的列
FLAG_FILL = PatternFill("solid", fgColor="FAEBDE")      # 优先核对的行
THIN = Side(style="thin", color="C8CFD4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INCLUDE_OPTS = ["是", "否", "存疑"]
EXCLUDE_OPTS = ["非纽扣电池", "磁性异物", "食管内定位", "重复就诊", "资料不全", "其他"]

# 列定义: (标题, 宽度, 是否换行, 是否填写列)
COLUMNS = [
    ("优先", 6, False, False),
    ("科研就诊编号", 14, False, False),
    ("年龄(岁)", 9, False, False),
    ("性别", 6, False, False),
    ("入院日期", 12, False, False),
    ("主诉", 22, True, False),
    ("现病史", 60, True, False),
    ("初步诊断", 20, True, False),
    ("出院诊断", 20, True, False),
    ("X线报告（时间｜结论｜自动判定）", 58, True, False),
    ("手术/内镜（名称｜术中诊断｜经过节选）", 58, True, False),
    ("自动:部位", 14, False, False),
    ("自动:处置", 18, False, False),
    ("自动:有症状", 10, False, False),
    ("自动:确认排出", 11, False, False),
    ("自动:需核对", 10, False, False),
    ("确认纳入", 10, False, True),
    ("排除原因", 14, False, True),
    ("核对备注", 26, True, True),
]


def clip(s, n):
    s = re.sub(r"\s+", " ", str(s)).strip()
    return s if len(s) <= n else s[:n] + "…"


# ============================================================
# 组装每例数据
# ============================================================
def build_rows():
    xl = C.load_workbook()
    base = xl.parse("病案首页基本信息")
    adm = xl.parse("儿科入院记录")
    xray = xl.parse("X线报告")
    op = xl.parse("住院病历手术记录")
    dis = xl.parse("住院病历出院记录")

    core = C.build_core_text(adm)
    battery = C.ascertain(core, C.BATTERY_KW)
    alltext = C.collect_text(xl, xl.sheet_names)

    xray = xray.copy()
    xray["_t"] = pd.to_datetime(xray["检查时间"], errors="coerce")
    bx = xray[xray["科研就诊编号"].isin(battery)].sort_values("_t")
    bo = op[op["科研就诊编号"].isin(battery)]
    adm_i = adm.set_index("科研就诊编号")
    dis_i = dis.drop_duplicates("科研就诊编号").set_index("科研就诊编号")
    base_i = base.set_index("科研就诊编号")

    rows = []
    for v in sorted(battery):
        b = base_i.loc[v]
        a = adm_i.loc[v] if v in adm_i.index else None
        d = dis_i.loc[v] if v in dis_i.index else None

        # ---- 影像：逐张列出，附自动判定 ----
        films, first_loc, any_absent = [], None, False
        g = bx[bx["科研就诊编号"] == v]
        for i, (_, r) in enumerate(g.iterrows()):
            txt = str(r["检查所见"]) + "。" + str(r["检查结论"])
            state, _ev = C.classify_film_state(txt)
            any_absent = any_absent or state == "absent"
            mark = {"present": "在位", "absent": "已排出", "unclear": "不明"}[state]
            t = r["_t"].strftime("%m-%d %H:%M") if pd.notna(r["_t"]) else "时间缺失"
            films.append(f"{i+1}) {t}｜{clip(r['检查结论'], 70)}｜[{mark}]")
            if i == 0:
                first_loc = C.classify_location(
                    str(r["检查所见"]) + " " + str(r["检查结论"]))
        film_txt = "\n".join(films) if films else "【无 X 线报告 — 需去 PACS 核查】"
        if not films:
            first_loc = "no_film"

        # ---- 操作 ----
        procs = []
        for _, r in bo[bo["科研就诊编号"] == v].iterrows():
            procs.append(
                f"{clip(r['手术名称'], 34)}｜{clip(r['术中诊断'], 40)}｜"
                f"{clip(r['手术经过'], 200)}"
            )
        proc_txt = "\n".join(procs) if procs else "（无手术/内镜记录 → 期待治疗）"
        proc_kinds = ";".join(sorted({
            C.classify_procedure(x) for x in
            bo[bo["科研就诊编号"] == v]["手术名称"]
        })) or "none_observation"

        flags = C.symptom_flags(core.get(v, ""))
        need_review = bool(
            any(k in alltext.get(v, "") for k in C.MAGNET_KW)
            or re.search(r"磁珠|巴克球", str(bo[bo["科研就诊编号"] == v]["术中诊断"].tolist()))
        )

        # ---- 优先级：越可能出问题越靠前 ----
        if need_review:
            pri = 1
        elif first_loc in ("no_film", "undetermined"):
            pri = 2
        elif first_loc == "abdomen_unspecified":
            pri = 3
        elif first_loc == "esophagus":
            pri = 4
        else:
            pri = 5

        rows.append({
            "优先": pri,
            "科研就诊编号": v,
            "年龄(岁)": b["年龄（岁）"],
            "性别": b["性别"],
            "入院日期": pd.to_datetime(b["入院日期"], errors="coerce").strftime("%Y-%m-%d"),
            "主诉": clip(a["主诉"], 40) if a is not None else "",
            "现病史": clip(a["现病史"], 420) if a is not None else "",
            "初步诊断": clip(a["初步诊断"], 60) if a is not None else "",
            "出院诊断": clip(d["出院诊断"], 60) if d is not None else "",
            "X线报告（时间｜结论｜自动判定）": film_txt,
            "手术/内镜（名称｜术中诊断｜经过节选）": proc_txt,
            "自动:部位": first_loc,
            "自动:处置": proc_kinds,
            "自动:有症状": "是" if any(flags.values()) else "否",
            "自动:确认排出": "是" if any_absent else "否",
            "自动:需核对": "★" if need_review else "",
            "确认纳入": "",
            "排除原因": "",
            "核对备注": "",
        })

    df = pd.DataFrame(rows).sort_values(
        ["优先", "科研就诊编号"]).reset_index(drop=True)
    return df


# ============================================================
# 写 Excel
# ============================================================
def write_xlsx(df, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    titles = [c[0] for c in COLUMNS]
    ws.append(titles)
    for j, (title, width, wrap, is_input) in enumerate(COLUMNS, start=1):
        col = get_column_letter(j)
        ws.column_dimensions[col].width = width
        c = ws.cell(row=1, column=j)
        c.font = Font(name=FONT, bold=True, size=10, color=INK)
        c.fill = HEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    for i, rec in enumerate(df.to_dict("records"), start=2):
        flagged = rec["自动:需核对"] == "★"
        for j, (title, width, wrap, is_input) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=i, column=j, value=rec[title])
            c.font = Font(name=FONT, size=9, color=INK)
            c.border = BORDER
            c.alignment = Alignment(
                vertical="top", wrap_text=wrap,
                horizontal="center" if not wrap and title != "现病史" else "left",
            )
            if is_input:
                c.fill = INPUT_FILL
            elif flagged:
                c.fill = FLAG_FILL
        ws.row_dimensions[i].height = 92

    n = len(df)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n + 1}"

    # 下拉选项
    ci = titles.index("确认纳入") + 1
    ei = titles.index("排除原因") + 1
    dv1 = DataValidation(type="list", formula1='"' + ",".join(INCLUDE_OPTS) + '"',
                         allow_blank=True, showDropDown=False)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(EXCLUDE_OPTS) + '"',
                         allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"{get_column_letter(ci)}2:{get_column_letter(ci)}{n + 1}")
    dv2.add(f"{get_column_letter(ei)}2:{get_column_letter(ei)}{n + 1}")

    # ---------------- 说明页 ----------------
    lg = wb.create_sheet(LEGEND)
    lg.column_dimensions["A"].width = 20
    lg.column_dimensions["B"].width = 88
    inc_col = get_column_letter(ci)

    lines = [
        ("纽扣电池队列　逐例人工核对工作表", None, True, 14),
        ("", None, False, 10),
        ("怎么用", None, True, 11),
        ("1", "打开「核对」页。每行一例，所有需要判断的病历原文已预先拼好。", False, 10),
        ("2", "只需填写三列黄色格：确认纳入 / 排除原因 / 核对备注。其余列请勿修改。", False, 10),
        ("3", "「优先」列已排好序：1 = 自动标记疑似假阳性（橙色行），2 = 无影像或部位未定，"
              "3 = 影像仅写腹部分区，4 = 食管内，5 = 常规。请从上往下做。", False, 10),
        ("4", "绝大多数病例读本表即可判断；仅在存疑时才需回 HIS 调阅原始病历。", False, 10),
        ("5", "填完后运行：python3 scripts/make_verification_sheet.py --to-csv "
              "output/核对工作表.xlsx　将生成 data/verified_cohort.csv，"
              "之后所有分析脚本自动改用核对后的队列。", False, 10),
        ("", None, False, 10),
        ("填写示例", None, True, 11),
        ("确认纳入", "是　　（确认为纽扣电池，纳入研究）", False, 10),
        ("确认纳入", "否　+　排除原因：磁性异物　+　备注：术中诊断为磁珠2颗，非电池", False, 10),
        ("确认纳入", "存疑　+　备注：现病史称可能为电池，影像未见双环征，待调阅原片", False, 10),
        ("", None, False, 10),
        ("务必注意", None, True, 11),
        ("已知假阳性", "关键词筛查已确证至少 1 例假阳性：病历文本提及「电池」，"
                      "但术中诊断实为磁珠（巴克球）致小肠穿孔。橙色行即为此类嫌疑。", False, 10),
        ("自动判定仅供参考", "「自动:」开头的各列是文本挖掘结果，不是结论。"
                          "与原文冲突时以原文为准，并在备注中写明。", False, 10),
        ("数据保密", "本表含患者层面病历文本，不得提交至代码仓库，不得对外传播。", False, 10),
        ("", None, False, 10),
        ("进度", None, True, 11),
    ]
    r = 1
    for a, b, bold, size in lines:
        lg.cell(row=r, column=1, value=a).font = Font(name=FONT, bold=bold, size=size, color=INK)
        if b is not None:
            c = lg.cell(row=r, column=2, value=b)
            c.font = Font(name=FONT, size=size, color=INK)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            lg.row_dimensions[r].height = 30 if len(b) > 60 else 16
        r += 1

    rng = f"'{SHEET}'!{inc_col}2:{inc_col}{n + 1}"
    for label, formula in [
        ("总例数", f"={n}"),
        ("已确认纳入", f'=COUNTIF({rng},"是")'),
        ("已排除", f'=COUNTIF({rng},"否")'),
        ("标记存疑", f'=COUNTIF({rng},"存疑")'),
        ("尚未处理", f'={n}-COUNTA({rng})'),
    ]:
        lg.cell(row=r, column=1, value=label).font = Font(name=FONT, size=10, color=INK)
        c = lg.cell(row=r, column=2, value=formula)
        c.font = Font(name=FONT, size=10, bold=True, color=INK)
        r += 1

    wb.save(path)
    return n


# ============================================================
# 转成最终队列
# ============================================================
def to_csv(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET]
    titles = [c.value for c in ws[1]]
    ci, ei, ni = (titles.index("确认纳入"), titles.index("排除原因"),
                  titles.index("科研就诊编号"))
    rows, unresolved = [], 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        inc = (r[ci] or "").strip() if isinstance(r[ci], str) else ""
        if inc == "":
            unresolved += 1
        rows.append({
            "科研就诊编号": r[ni],
            "include": 1 if inc == "是" else 0,
            "verify_status": inc or "未处理",
            "exclude_reason": r[ei] or "",
        })
    df = pd.DataFrame(rows)
    os.makedirs("data", exist_ok=True)
    out = "data/verified_cohort.csv"
    df.to_csv(out, index=False, encoding="utf-8-sig")
    kept = int(df["include"].sum())
    print(f"已写出 {out}")
    print(f"  纳入 {kept} 例 / 排除 {int((df['verify_status']=='否').sum())} 例 / "
          f"存疑 {int((df['verify_status']=='存疑').sum())} 例 / 未处理 {unresolved} 例")
    if unresolved:
        print(f"  !! 仍有 {unresolved} 例未填写「确认纳入」，这些例已按未纳入处理")
    if df["verify_status"].eq("存疑").any():
        print("  !! 存疑病例需裁定后再改为是/否，否则不计入队列")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--to-csv", metavar="XLSX", default=None,
                    help="把填好的工作表转成 data/verified_cohort.csv")
    args = ap.parse_args()

    if args.to_csv:
        to_csv(args.to_csv)
        return

    df = build_rows()
    n = write_xlsx(df, OUTFILE)
    print(f"已生成 {OUTFILE}（{n} 例）")
    print("优先级分布:", df["优先"].value_counts().sort_index().to_dict())
    print("  1=疑似假阳性  2=无影像/部位未定  3=仅腹部分区  4=食管内  5=常规")
    print(f"\n填完后运行: python3 scripts/make_verification_sheet.py --to-csv {OUTFILE}")


if __name__ == "__main__":
    main()
